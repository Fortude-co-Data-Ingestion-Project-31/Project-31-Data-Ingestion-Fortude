"""Read supported documents from a SharePoint document library via Graph."""

import os
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

import requests
from docx import Document
from dotenv import load_dotenv
from pypdf import PdfReader


GRAPH_BASE_URL = "https://graph.microsoft.com/v1.0"
SUPPORTED_FILE_TYPES = {".txt", ".docx", ".pdf", ".pptx", ".xlsx"}


class SharePointDeltaStateError(RuntimeError):
    """Raised when Graph no longer accepts a saved delta cursor."""


def _required_setting(name):
    """Read a required SharePoint setting and fail clearly when it is missing."""
    value = os.getenv(name)
    if not value:
        raise RuntimeError(f"Missing required environment variable: {name}")
    return value


def _get_access_token(session, tenant_id, client_id, client_secret):
    """Authenticate the app with Microsoft and return a Graph access token."""
    response = session.post(
        f"https://login.microsoftonline.com/{tenant_id}/oauth2/v2.0/token",
        data={
            "client_id": client_id,
            "client_secret": client_secret,
            "scope": "https://graph.microsoft.com/.default",
            "grant_type": "client_credentials",
        },
        timeout=30,
    )
    response.raise_for_status()
    return response.json()["access_token"]


def _get_json(session, url, headers):
    """Request one Graph response and return its JSON body."""
    response = session.get(url, headers=headers, timeout=30)
    response.raise_for_status()
    return response.json()


def _get_all_items(session, url, headers):
    """Follow Graph pagination links and combine all returned items."""
    items = []
    while url:
        page = _get_json(session, url, headers)
        items.extend(page.get("value", []))
        url = page.get("@odata.nextLink")
    return items


def _extract_document(file_name, file_bytes):
    """Extract text and location-aware units from a supported document."""
    file_type = Path(file_name).suffix.lower()

    if file_type == ".txt":
        content = file_bytes.decode("utf-8-sig")
        return content, [{"content": content}]
    if file_type == ".docx":
        document = Document(BytesIO(file_bytes))
        paragraphs = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                paragraphs.append({"content": text})
        return "\n".join(item["content"] for item in paragraphs), paragraphs
    if file_type == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        pages = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append({"content": text, "page": page_number})
        return "\n".join(page["content"] for page in pages), pages
    if file_type == ".pptx":
        from pptx import Presentation

        presentation = Presentation(BytesIO(file_bytes))
        slides = []
        # Keep slide numbers so chunks can point back to their source slide.
        for slide_number, slide in enumerate(presentation.slides, start=1):
            text_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_parts.append(shape.text.strip())
            if text_parts:
                slides.append({"content": "\n".join(text_parts), "slide": slide_number})
        return "\n".join(slide["content"] for slide in slides), slides
    if file_type == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
        rows = []
        try:
            for worksheet in workbook.worksheets:
                # Keep the sheet name so spreadsheet chunks retain useful context.
                for row in worksheet.iter_rows(values_only=True):
                    values = []
                    for value in row:
                        if value is not None and str(value).strip():
                            values.append(str(value))
                    if values:
                        rows.append({"content": " | ".join(values), "sheet": worksheet.title})
        finally:
            workbook.close()
        return "\n".join(row["content"] for row in rows), rows

    raise ValueError(f"Unsupported file type: {file_type}")


def _extract_content(file_name, file_bytes):
    """Return combined text for compatibility with existing callers."""
    content, _ = _extract_document(file_name, file_bytes)
    return content


def _author_from_item(item):
    """Return the best available author name from a Graph drive item."""
    identity = item.get("lastModifiedBy") or item.get("createdBy") or {}
    person = identity.get("user") or identity.get("application") or {}
    return person.get("displayName") or person.get("email") or person.get("id")


def _get_view_count(session, drive_id, item_id, headers):
    """Return Graph's all-time access count, or None when analytics is unavailable."""
    try:
        analytics = _get_json(
            session,
            f"{GRAPH_BASE_URL}/drives/{drive_id}/items/{item_id}/analytics",
            headers,
        )
        return analytics["allTime"]["access"]["actionCount"]
    except (requests.RequestException, KeyError, TypeError):
        return None


def _resolve_drive(session):
    """Authenticate and resolve the configured site, library, and root folder."""
    load_dotenv()
    tenant_id = _required_setting("TENANT_ID")
    client_id = _required_setting("CLIENT_ID")
    client_secret = _required_setting("CLIENT_SECRET")
    host = _required_setting("SHAREPOINT_HOST").strip().rstrip("/")
    site_path = _required_setting("SHAREPOINT_SITE_PATH").strip().strip("/")
    library_name = _required_setting("SHAREPOINT_LIBRARY").strip()

    token = _get_access_token(session, tenant_id, client_id, client_secret)
    headers = {"Authorization": f"Bearer {token}"}

    encoded_site_path = quote(site_path, safe="/")
    site = _get_json(
        session,
        f"{GRAPH_BASE_URL}/sites/{host}:/{encoded_site_path}",
        headers,
    )
    drives = _get_all_items(
        session, f"{GRAPH_BASE_URL}/sites/{site['id']}/drives", headers
    )
    drive = next(
        (item for item in drives if item.get("name", "").casefold() == library_name.casefold()),
        None,
    )
    if drive is None:
        raise RuntimeError(f"SharePoint document library not found: {library_name}")

    root = _get_json(
        session, f"{GRAPH_BASE_URL}/drives/{drive['id']}/root", headers
    )
    return drive["id"], root["id"], headers


def get_sharepoint_drive_id(session=None):
    """Resolve and return the configured SharePoint document-library drive ID."""
    session = session or requests.Session()
    drive_id, _, _ = _resolve_drive(session)
    return drive_id


def _download_file_record(session, drive_id, headers, item):
    """Download one drive item and build the raw record used by the mapper."""
    file_name = item.get("name", "")
    file_type = Path(file_name).suffix.lower()
    download_response = session.get(
        f"{GRAPH_BASE_URL}/drives/{drive_id}/items/{item['id']}/content",
        headers=headers,
        timeout=60,
    )
    download_response.raise_for_status()
    content, content_units = _extract_document(file_name, download_response.content)
    return {
        "source": "sharepoint",
        "file_name": file_name,
        "file_type": file_type,
        "file_size": item.get("size"),
        "modified_at": item.get("lastModifiedDateTime"),
        "content": content,
        "content_units": content_units,
        "source_url": item.get("webUrl"),
        "author": _author_from_item(item),
        "version": item.get("eTag") or item.get("cTag"),
        "view_count": _get_view_count(session, drive_id, item["id"], headers),
    }


def read_sharepoint_delta(delta_url=None, session=None):
    """Read direct-root changes and return a cursor for the next delta sync."""
    session = session or requests.Session()
    drive_id, root_id, headers = _resolve_drive(session)
    url = delta_url or f"{GRAPH_BASE_URL}/drives/{drive_id}/root/delta"
    latest_by_id = {}
    final_delta_link = None

    while url:
        response = session.get(url, headers=headers, timeout=30)
        if response.status_code in (404, 410) and delta_url:
            raise SharePointDeltaStateError("Saved SharePoint delta cursor is invalid.")
        response.raise_for_status()
        page = response.json()
        # Later entries for the same item describe its final state in this sync.
        for item in page.get("value", []):
            item_id = item.get("id")
            if item_id:
                latest_by_id[item_id] = item
        url = page.get("@odata.nextLink")
        if not url:
            final_delta_link = page.get("@odata.deltaLink")

    if not final_delta_link:
        raise RuntimeError("Microsoft Graph delta response did not include @odata.deltaLink.")

    changes = []
    for item_id, item in latest_by_id.items():
        deleted = "deleted" in item
        file_name = item.get("name", "")
        file_type = Path(file_name).suffix.lower()
        # This feature only ingests supported files placed directly in the library root.
        is_direct_root_file = (
            "file" in item
            and (item.get("parentReference") or {}).get("id") == root_id
        )
        supported = is_direct_root_file and file_type in SUPPORTED_FILE_TYPES
        change = {
            "item_id": item_id,
            "deleted": deleted,
            "supported": supported,
            "record": None,
        }
        if supported and not deleted:
            change["record"] = _download_file_record(
                session, drive_id, headers, item
            )
        changes.append(change)

    return {
        "changes": changes,
        "delta_link": final_delta_link,
    }
